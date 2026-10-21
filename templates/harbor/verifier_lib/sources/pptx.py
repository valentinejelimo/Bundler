from pathlib import Path

from pydantic import field_validator

from ..models import StrictModel
from ..source_types import SourceCommand, SourceContext
from ._content_limits import finalize_accumulated


class PptxExtractTextInput(StrictModel):
    """Input for pptx.extract_text.

    Attributes:
        path: Workspace-relative PPTX path.
    """

    path: str

    @field_validator("path")
    @classmethod
    def require_pptx_extension(cls, value: str) -> str:
        """Requires the path to target a PPTX file.

        Args:
            value: Workspace-relative path from verifier.json.

        Returns:
            The unchanged path when it has a PPTX extension.

        Raises:
            ValueError: If the path does not end in ``.pptx``.
        """
        if Path(value).suffix.lower() != ".pptx":
            raise ValueError("pptx.extract_text requires a .pptx file")
        return value


class PptxExtractTextOutput(StrictModel):
    """Output from pptx.extract_text.

    Attributes:
        text: Extracted slide text.
    """

    text: str


class ExtractText(SourceCommand):
    """Extracts text from PPTX slides."""

    name = "extract_text"
    input_model = PptxExtractTextInput
    output_model = PptxExtractTextOutput

    def run(self, source_input: PptxExtractTextInput, context: SourceContext) -> PptxExtractTextOutput:
        """Runs PPTX text extraction.

        Args:
            source_input: Validated PPTX extraction input.
            context: Source runtime context.

        Returns:
            Extracted PPTX text.
        """
        resolved = context.resolve_path(source_input.path)
        return PptxExtractTextOutput(text=extract_pptx_text(resolved, context.max_content_chars))


def extract_pptx_text(path: Path, limit: int) -> str:
    """Extracts visible slide and table text from a PPTX file.

    Args:
        path: Resolved PPTX file path.
        limit: Maximum number of characters to return.

    Returns:
        Extracted text capped to ``limit`` characters. If the deck is longer than the cap, a
        truncation notice is appended so the cut is visible to the grader instead of silent.
    """
    from pptx import Presentation

    presentation = Presentation(path)
    parts: list[str] = []
    total = 0
    truncated = False
    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts = list(iter_slide_text(slide))
        if not slide_parts:
            continue
        total = append_limited(parts, f"[Slide {index}]", total, limit)
        if total >= limit:
            truncated = True
            break
        for text in slide_parts:
            total = append_limited(parts, text, total, limit)
            if total >= limit:
                truncated = True
                break
        if truncated:
            break
    return finalize_accumulated(parts, limit, truncated)


def iter_slide_text(slide: object) -> list[str]:
    """Collects text from a slide's text frames and tables.

    Args:
        slide: python-pptx slide object.

    Returns:
        Non-empty text fragments from the slide.
    """
    fragments: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text.strip()
            if text:
                fragments.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    fragments.append(" | ".join(cells))
    return fragments


def append_limited(parts: list[str], text: str, total: int, limit: int) -> int:
    """Appends text and returns the updated character count.

    Args:
        parts: Text fragments collected so far.
        text: Text fragment to append.
        total: Current approximate character count.
        limit: Maximum desired character count.

    Returns:
        Updated approximate character count.
    """
    parts.append(text)
    return total + len(text) + 1


COMMANDS = (ExtractText(),)
